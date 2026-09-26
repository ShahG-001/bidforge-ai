from io import BytesIO


def read_uploaded_file(uploaded_file):
    """Extract text from PDF, DOCX, TXT, Markdown, or CSV upload."""
    extension = uploaded_file.name.lower().rsplit(".", 1)[-1]
    try:
        data = uploaded_file.getvalue()
        if extension in {"txt", "md", "csv"}:
            text = data.decode("utf-8", errors="replace")
            if extension == "csv":
                lines = text.splitlines()
                text = "\n".join(f"[SOURCE: {uploaded_file.name} | ROW {index}] {line}" for index, line in enumerate(lines, start=1))
            return f"[SOURCE: {uploaded_file.name}]\n{text}", ""
        if extension == "pdf":
            from pypdf import PdfReader

            pages = PdfReader(BytesIO(data)).pages
            page_texts = [page.extract_text() or "" for page in pages]
            if not any(page_text.strip() for page_text in page_texts):
                try:
                    from pdf2image import convert_from_bytes
                    import pytesseract

                    images = convert_from_bytes(data, dpi=180, first_page=1, last_page=min(len(pages), 30))
                    text = "\n\n".join(f"[SOURCE: {uploaded_file.name} | PAGE {index} | OCR]\n{pytesseract.image_to_string(image)}" for index, image in enumerate(images, start=1))
                    if not text.strip():
                        return "", f"OCR found no readable text in {uploaded_file.name}. Check the scan or paste the relevant text."
                    return text, f"{uploaded_file.name} was scanned with OCR; verify the extracted text carefully. OCR is limited to the first 30 pages."
                except Exception as ocr_error:
                    return "", f"{uploaded_file.name} appears scanned and OCR could not read it ({ocr_error}). Upload a searchable PDF or paste its text."
            text = "\n\n".join(f"[SOURCE: {uploaded_file.name} | PAGE {index}]\n{page_text}" for index, page_text in enumerate(page_texts, start=1))
            return text, ""
        if extension == "docx":
            from docx import Document

            document = Document(BytesIO(data))
            chunks = [f"[SOURCE: {uploaded_file.name} | PARAGRAPH {index}] {p.text}" for index, p in enumerate(document.paragraphs, start=1) if p.text.strip()]
            for table_index, table in enumerate(document.tables, start=1):
                chunks.append(f"[SOURCE: {uploaded_file.name} | TABLE {table_index}]")
                chunks.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
            return "\n".join(chunks), ""
        return "", f"Unsupported file: {uploaded_file.name}. Please use PDF, DOCX, TXT, MD, or CSV."
    except Exception as error:
        return "", f"Could not read {uploaded_file.name}: {error}"
