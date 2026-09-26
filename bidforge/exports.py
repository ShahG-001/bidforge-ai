"""Render generated Markdown into clean Word and PDF documents."""
import re
from io import BytesIO


def _cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_separator(line: str) -> bool:
    cells = _cells(line)
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", cell.replace(" ", "")) for cell in cells)


def _plain(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\*\*(.*?)\*\*|__(.*?)__", lambda m: m.group(1) or m.group(2), text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)|(?<!_)_([^_]+)_(?!_)", lambda m: m.group(1) or m.group(2), text)
    return text.replace("`", "").strip()


def _markdown_lines(markdown_text: str):
    """Yield paragraph/table blocks while keeping Markdown table rows together."""
    lines = markdown_text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                if not _is_separator(lines[index]):
                    table_lines.append(_cells(lines[index]))
                index += 1
            if table_lines:
                width = max(len(row) for row in table_lines)
                yield "table", [row + [""] * (width - len(row)) for row in table_lines]
            continue
        yield "text", line
        index += 1


def to_docx(markdown_text: str) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(31, 41, 55)
    doc.add_heading("BidForge AI — Tender Response Draft", 0)

    for kind, content in _markdown_lines(markdown_text):
        if kind == "table":
            rows = content
            table = doc.add_table(rows=1, cols=len(rows[0]))
            table.style = "Table Grid"
            table.autofit = True
            for row_index, row_data in enumerate(rows):
                cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
                for col_index, value in enumerate(row_data):
                    cell = cells[col_index]
                    cell.text = _plain(value)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(2)
                        for run in paragraph.runs:
                            run.font.size = Pt(8)
                            if row_index == 0:
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
                    if row_index == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "1D4ED8")
                        cell._tc.get_or_add_tcPr().append(shading)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        stripped = content
        if not stripped:
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            doc.add_heading(_plain(heading.group(2)), level=level)
        elif re.match(r"^\s*[-*+]\s+", stripped):
            doc.add_paragraph(_plain(re.sub(r"^\s*[-*+]\s+", "", stripped)), style="List Bullet")
        elif re.match(r"^\s*\d+[.)]\s+", stripped):
            doc.add_paragraph(_plain(re.sub(r"^\s*\d+[.)]\s+", "", stripped)), style="List Number")
        else:
            paragraph = doc.add_paragraph(_plain(stripped))
            paragraph.paragraph_format.space_after = Pt(5)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def to_pdf(markdown_text: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    output = BytesIO()
    page_width, _ = A4
    margin = 17 * mm
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=margin, leftMargin=margin, topMargin=18 * mm, bottomMargin=18 * mm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="BFBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.8, leading=12.3, spaceAfter=5, textColor=colors.HexColor("#1f2937"), alignment=TA_LEFT, wordWrap="CJK"))
    styles.add(ParagraphStyle(name="BFTable", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.4, leading=9.2, textColor=colors.HexColor("#1f2937"), wordWrap="CJK"))
    styles.add(ParagraphStyle(name="BFTableHead", parent=styles["BFTable"], fontName="Helvetica-Bold", textColor=colors.white))
    story = [Paragraph("BidForge AI — Tender Response Draft", styles["Title"]), Spacer(1, 8)]
    available_width = page_width - 2 * margin

    for kind, content in _markdown_lines(markdown_text):
        if kind == "table":
            rows = content
            table_data = []
            for row_index, row in enumerate(rows):
                style = styles["BFTableHead"] if row_index == 0 else styles["BFTable"]
                table_data.append([Paragraph(escape(_plain(cell)), style) for cell in row])
            widths = [available_width / len(rows[0])] * len(rows[0])
            table = Table(table_data, colWidths=widths, repeatRows=1, hAlign="LEFT", splitByRow=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1d4ed8")),
                ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#cbd5e1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f8fc")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.extend([table, Spacer(1, 8)])
            continue

        stripped = content
        if not stripped:
            story.append(Spacer(1, 3))
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            style = styles["Heading1"] if level == 1 else styles["Heading2"] if level == 2 else styles["Heading3"]
            story.append(Paragraph(escape(_plain(heading.group(2))), style))
        elif re.match(r"^\s*[-*+]\s+", stripped):
            item = _plain(re.sub(r"^\s*[-*+]\s+", "", stripped))
            story.append(Paragraph("&#8226;&nbsp; " + escape(item), styles["BFBody"]))
        elif re.match(r"^\s*\d+[.)]\s+", stripped):
            item = _plain(stripped)
            story.append(Paragraph(escape(item), styles["BFBody"]))
        else:
            story.append(Paragraph(escape(_plain(stripped)), styles["BFBody"]))

    doc.build(story)
    return output.getvalue()
